from fastapi import FastAPI, HTTPException, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from typing import List, Dict, Any
import pandas as pd
import io
import time
from datetime import datetime

app = FastAPI(title="表格处理API", description="处理多种格式文件的API服务")

# 配置 CORS
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],  # 允许所有来源，生产环境建议指定具体域名
    allow_credentials=True,
    allow_methods=["*"],  # 允许所有方法
    allow_headers=["*"],  # 允许所有请求头
)

# 挂载 ML 路由
from ml import ml_router
app.include_router(ml_router)


@app.on_event("startup")
async def on_startup():
    import asyncio
    from ml.cleanup import cleanup_expired_sessions
    asyncio.create_task(cleanup_expired_sessions())


def process_tables(tables: List[pd.DataFrame]) -> Dict:
    """处理多个表格数据"""
    result = {
        'status': 'success',
        'message': '',
        'processed_count': 0,
        'table_info': []
    }
    
    try:
        for i, df in enumerate(tables):
            if df.empty:
                table_info = {
                    'index': i,
                    'shape': (0, 0),
                    'columns': [],
                    'sample_data': [],
                    'is_valid': True
                }
            else:
                table_info = {
                    'index': i,
                    'shape': df.shape,
                    'columns': list(df.columns),
                    'sample_data': df.head(3).to_dict('records'),
                    'is_valid': True
                }
            
            result['table_info'].append(table_info)
        
        result['processed_count'] = len(tables)
        result['message'] = f'成功处理 {len(tables)} 个表格'
        return result
        
    except Exception as e:
        return {
            'status': 'failure',
            'message': f'处理表格时发生错误: {str(e)}',
            'processed_count': 0,
            'table_info': []
        }


def parse_file(filename: str, content: bytes) -> List[pd.DataFrame]:
    """根据文件类型解析文件，返回 DataFrame 列表。支持 CSV、Excel(xls/xlsx)、TXT、JSON 等"""
    filename_lower = filename.lower()
    buf = io.BytesIO(content)

    # CSV
    if filename_lower.endswith('.csv'):
        return [pd.read_csv(buf)]

    # Excel: .xls 用 xlrd，.xlsx 用 openpyxl
    if filename_lower.endswith('.xlsx'):
        excel_file = pd.ExcelFile(buf, engine='openpyxl')
        return [pd.read_excel(excel_file, sheet_name=sheet) for sheet in excel_file.sheet_names]
    if filename_lower.endswith('.xls'):
        excel_file = pd.ExcelFile(buf, engine='xlrd')
        return [pd.read_excel(excel_file, sheet_name=sheet) for sheet in excel_file.sheet_names]

    # TXT（制表符或逗号分隔）
    if filename_lower.endswith('.txt'):
        text = content.decode('utf-8', errors='replace')
        if '\t' in text.split('\n')[0]:
            return [pd.read_csv(io.BytesIO(content), sep='\t')]
        return [pd.read_csv(io.BytesIO(content))]

    # JSON（支持数组或键值结构）
    if filename_lower.endswith('.json'):
        try:
            return [pd.read_json(buf)]
        except Exception:
            buf.seek(0)
            try:
                return [pd.read_json(buf, orient='records')]
            except Exception:
                buf.seek(0)
                import json
                data = json.load(buf)
                if isinstance(data, list):
                    return [pd.DataFrame(data)]
                if isinstance(data, dict):
                    return [pd.DataFrame([data])]
                return [pd.DataFrame({'data': [str(data)]})]
    
    # Word (docx) - 提取表格
    elif filename_lower.endswith('.docx'):
        from docx import Document
        doc = Document(io.BytesIO(content))
        dataframes = []
        for table in doc.tables:
            data = []
            for row in table.rows:
                data.append([cell.text for cell in row.cells])
            if data:
                df = pd.DataFrame(data[1:], columns=data[0]) if len(data) > 1 else pd.DataFrame(data)
                dataframes.append(df)
        if not dataframes:
            # 如果没有表格，将段落文本作为单列DataFrame
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            dataframes = [pd.DataFrame({'content': paragraphs})]
        return dataframes
    
    # PDF - 提取表格
    elif filename_lower.endswith('.pdf'):
        import pdfplumber
        dataframes = []
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            for page in pdf.pages:
                tables = page.extract_tables()
                for table in tables:
                    if table:
                        df = pd.DataFrame(table[1:], columns=table[0]) if len(table) > 1 else pd.DataFrame(table)
                        dataframes.append(df)
                # 如果没有表格，提取文本
                if not tables:
                    text = page.extract_text()
                    if text:
                        dataframes.append(pd.DataFrame({'content': [text]}))
        return dataframes if dataframes else [pd.DataFrame({'content': ['PDF无内容']})]
    
    else:
        raise ValueError(f"不支持的文件格式: {filename}")


@app.post("/upload")
async def upload_files(files: List[UploadFile] = File(default=[])):
    """
    统一文件上传接口（支持多文件）
    
    支持格式: CSV, Excel(xlsx/xls), TXT, JSON, Word(docx), PDF
    """
    start_time = time.time()
    request_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    
    if not files:
        raise HTTPException(status_code=400, detail="请上传至少一个文件")
    
    try:
        all_dataframes = []
        file_results = []
        process_logs = []
        
        process_logs.append({
            'step': '开始处理',
            'message': f'收到 {len(files)} 个文件',
            'timestamp': datetime.now().strftime("%H:%M:%S.%f")[:-3]
        })
        
        for idx, file in enumerate(files):
            file_start_time = time.time()
            file_size = 0
            
            try:
                content = await file.read()
                file_size = len(content)
                process_logs.append({
                    'step': f'文件 {idx + 1}: {file.filename}',
                    'message': f'文件大小: {file_size / 1024:.2f} KB',
                    'timestamp': datetime.now().strftime("%H:%M:%S.%f")[:-3]
                })
                
                parse_start = time.time()
                dfs = parse_file(file.filename, content)
                parse_time = time.time() - parse_start
                
                all_dataframes.extend(dfs)
                file_time = time.time() - file_start_time
                
                file_results.append({
                    'filename': file.filename,
                    'status': 'success',
                    'tables_found': len(dfs),
                    'file_size_kb': round(file_size / 1024, 2),
                    'parse_time_ms': round(parse_time * 1000, 2),
                    'total_time_ms': round(file_time * 1000, 2)
                })
                
                process_logs.append({
                    'step': f'文件 {idx + 1}: {file.filename}',
                    'message': f'解析成功，找到 {len(dfs)} 个表格，耗时 {parse_time * 1000:.2f}ms',
                    'timestamp': datetime.now().strftime("%H:%M:%S.%f")[:-3]
                })
                
            except Exception as e:
                file_time = time.time() - file_start_time
                file_results.append({
                    'filename': file.filename,
                    'status': 'error',
                    'error': str(e),
                    'file_size_kb': round(file_size / 1024, 2) if file_size > 0 else 0,
                    'total_time_ms': round(file_time * 1000, 2)
                })
                
                process_logs.append({
                    'step': f'文件 {idx + 1}: {file.filename}',
                    'message': f'处理失败: {str(e)}',
                    'timestamp': datetime.now().strftime("%H:%M:%S.%f")[:-3],
                    'level': 'error'
                })
        
        process_start = time.time()
        result = process_tables(all_dataframes)
        process_time = time.time() - process_start
        
        total_time = time.time() - start_time
        
        result['file_results'] = file_results
        result['runtime_info'] = {
            'request_time': request_time,
            'total_files': len(files),
            'success_files': len([f for f in file_results if f['status'] == 'success']),
            'failed_files': len([f for f in file_results if f['status'] == 'error']),
            'total_tables': result['processed_count'],
            'total_time_ms': round(total_time * 1000, 2),
            'process_time_ms': round(process_time * 1000, 2),
            'process_logs': process_logs
        }
        
        process_logs.append({
            'step': '处理完成',
            'message': f'总共处理 {result["processed_count"]} 个表格，总耗时 {total_time * 1000:.2f}ms',
            'timestamp': datetime.now().strftime("%H:%M:%S.%f")[:-3]
        })
        
        return result
    
    except Exception as e:
        total_time = time.time() - start_time
        raise HTTPException(
            status_code=500, 
            detail={
                'error': str(e),
                'request_time': request_time,
                'total_time_ms': round(total_time * 1000, 2)
            }
        )


@app.post("/upload-single")
async def upload_single_file(file: UploadFile = File(...)):
    """
    单文件上传接口
    
    支持格式: CSV, Excel(xlsx/xls), TXT, JSON, Word(docx), PDF
    
    注意：参数名必须是 'file'，使用 form-data 格式上传
    """
    start_time = time.time()
    request_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    
    if not file:
        raise HTTPException(status_code=400, detail="请上传文件，参数名必须是 'file'")
    
    try:
        process_logs = []
        
        process_logs.append({
            'step': '开始处理',
            'message': f'收到文件: {file.filename}',
            'timestamp': datetime.now().strftime("%H:%M:%S.%f")[:-3]
        })
        
        read_start = time.time()
        content = await file.read()
        file_size = len(content)
        read_time = time.time() - read_start
        
        process_logs.append({
            'step': '读取文件',
            'message': f'文件大小: {file_size / 1024:.2f} KB，读取耗时: {read_time * 1000:.2f}ms',
            'timestamp': datetime.now().strftime("%H:%M:%S.%f")[:-3]
        })
        
        parse_start = time.time()
        dfs = parse_file(file.filename, content)
        parse_time = time.time() - parse_start
        
        process_logs.append({
            'step': '解析文件',
            'message': f'解析成功，找到 {len(dfs)} 个表格，耗时: {parse_time * 1000:.2f}ms',
            'timestamp': datetime.now().strftime("%H:%M:%S.%f")[:-3]
        })
        
        process_start = time.time()
        result = process_tables(dfs)
        process_time = time.time() - process_start
        
        total_time = time.time() - start_time
        
        result['filename'] = file.filename
        result['runtime_info'] = {
            'request_time': request_time,
            'filename': file.filename,
            'file_size_kb': round(file_size / 1024, 2),
            'tables_found': len(dfs),
            'read_time_ms': round(read_time * 1000, 2),
            'parse_time_ms': round(parse_time * 1000, 2),
            'process_time_ms': round(process_time * 1000, 2),
            'total_time_ms': round(total_time * 1000, 2),
            'process_logs': process_logs
        }
        
        process_logs.append({
            'step': '处理完成',
            'message': f'总耗时: {total_time * 1000:.2f}ms',
            'timestamp': datetime.now().strftime("%H:%M:%S.%f")[:-3]
        })
        
        return result
    except Exception as e:
        total_time = time.time() - start_time
        raise HTTPException(
            status_code=500, 
            detail={
                'error': str(e),
                'filename': file.filename,
                'request_time': request_time,
                'total_time_ms': round(total_time * 1000, 2)
            }
        )


@app.get("/")
async def root():
    return {
        "message": "文件处理API服务",
        "supported_formats": ["csv", "xlsx", "xls", "txt", "json", "docx", "pdf"],
        "docs": "/docs"
    }


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)
